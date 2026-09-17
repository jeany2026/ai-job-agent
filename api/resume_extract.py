"""Extract plain text from an uploaded resume. No semantic analysis."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

MAX_RESUME_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc"}
ATTACHMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md"}
ALLOWED_ATTACHMENT_KINDS = {
    "resume",
    "project_document",
    "other_attachment",
    "user_statement",
}
DOC_UNSUPPORTED_MESSAGE = "当前环境暂不具备 DOC 文件解析能力，请转换为 DOCX 或 PDF。"
ATTACHMENT_UNREADABLE_MESSAGE = "附件无法解析，请换成可复制文本的 PDF、DOCX 或 TXT。"


class ResumeExtractError(ValueError):
    def __init__(self, error_code: str, message: str):
        super().__init__(message)
        self.error_code = error_code
        self.message = message


def validate_resume_file(*, filename: str | None, data: bytes | None) -> str:
    return _validate_upload(
        filename=filename,
        data=data,
        allowed=ALLOWED_EXTENSIONS,
        empty_message="简历文件是空的。",
        oversize_message="简历文件不能超过 10MB。",
        unsupported_message="仅支持 PDF / DOC / DOCX。",
    )


def _validate_upload(
    *,
    filename: str | None,
    data: bytes | None,
    allowed: set[str],
    empty_message: str,
    oversize_message: str,
    unsupported_message: str,
    missing_code: str = "empty_file",
    missing_message: str | None = None,
) -> str:
    if data is None:
        raise ResumeExtractError(missing_code, missing_message or empty_message)
    name = _safe_filename(filename)
    if not name:
        raise ResumeExtractError(missing_code, missing_message or empty_message)
    suffix = Path(name).suffix.lower()
    if suffix not in allowed:
        raise ResumeExtractError("unsupported_extension", unsupported_message)
    if len(data) == 0:
        raise ResumeExtractError("empty_file", empty_message)
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeExtractError("file_too_large", oversize_message)
    return suffix


def normalize_attachment_kind(kind: str | None) -> str:
    text = (kind or "").strip()
    if not text:
        return "other_attachment"
    if text not in ALLOWED_ATTACHMENT_KINDS:
        return "other_attachment"
    return text


def extract_attachment_text(*, filename: str | None, data: bytes | None) -> str:
    suffix = _validate_upload(
        filename=filename,
        data=data,
        allowed=ATTACHMENT_EXTENSIONS,
        empty_message="附件是空的。",
        oversize_message="附件不能超过 10MB。",
        unsupported_message="附件仅支持 PDF / DOC / DOCX / TXT / MD。",
    )
    assert data is not None
    if suffix == ".doc":
        raise ResumeExtractError("doc_unsupported", DOC_UNSUPPORTED_MESSAGE)
    if suffix in {".txt", ".md"}:
        text = _extract_plain_text(data)
    else:
        text = _extract_office_file(data=data, suffix=suffix, error_message=ATTACHMENT_UNREADABLE_MESSAGE)
    cleaned = (text or "").strip()
    if not cleaned:
        raise ResumeExtractError("resume_unreadable", ATTACHMENT_UNREADABLE_MESSAGE)
    return cleaned


def extract_resume_text(*, filename: str | None, data: bytes | None) -> str:
    suffix = validate_resume_file(filename=filename, data=data)
    assert data is not None
    if suffix == ".doc":
        raise ResumeExtractError("doc_unsupported", DOC_UNSUPPORTED_MESSAGE)

    resume_unreadable = "简历无法解析，请换成可复制文本的 PDF 或 DOCX。"
    text = _extract_office_file(data=data, suffix=suffix, error_message=resume_unreadable)
    cleaned = (text or "").strip()
    if not cleaned:
        raise ResumeExtractError("resume_unreadable", resume_unreadable)
    return cleaned


def _extract_office_file(*, data: bytes, suffix: str, error_message: str) -> str:
    tmp_path = None
    try:
        fd, tmp_path = tempfile.mkstemp(prefix="aijob-upload-", suffix=suffix)
        os.close(fd)
        Path(tmp_path).write_bytes(data)
        if suffix == ".pdf":
            return _extract_pdf(tmp_path)
        return _extract_docx(tmp_path)
    except ResumeExtractError:
        raise
    except Exception:
        raise ResumeExtractError("resume_unreadable", error_message) from None
    finally:
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def _extract_plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeExtractError("resume_unreadable", ATTACHMENT_UNREADABLE_MESSAGE)


def _safe_filename(filename: str | None) -> str:
    return Path(filename or "").name.strip()


def _extract_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    chunks: list[str] = []
    for page in reader.pages:
        chunk = page.extract_text() or ""
        if chunk.strip():
            chunks.append(chunk)
    return "\n".join(chunks)


def _extract_docx(path: str) -> str:
    from docx import Document

    document = Document(path)
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                lines.append(" ".join(cells))
    return "\n".join(lines)
